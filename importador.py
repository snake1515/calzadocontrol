"""Importación de datos desde Excel.

La empresa descarga una plantilla con una hoja por sección, la llena (copiando y pegando
desde sus propios archivos) y la sube. La app primero hace un ensayo completo dentro de una
transacción que se deshace (vista previa con errores fila por fila) y solo si todo está bien
la aplica de verdad. Cada importación queda como un lote que se puede deshacer.
"""
import io
import json
import re
import unicodedata
from datetime import date, datetime, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from models import (CATEGORIAS_GASTO, CATEGORIAS_INSUMO, LINEAS, PROCESOS, TIPOS_CLIENTE, UNIDADES,
                    Activo, Cliente, CompraInsumo, Empleado, Gasto, Insumo, PagoNomina, PrepRecetaItem,
                    ProcesoTarifa, Produccion, ProduccionConsumo, ProduccionMO, Producto, RecetaItem,
                    Venta, VentaItem, ahora, db)
from services import ErrorNegocio, anular_produccion, promedio_ponderado, registrar_produccion
from unidades import FACTORES, convertir

TALLAS = list(range(21, 46))
MAX_FILAS = 20000
MARCA_AYUDA = "▸"

SI_NO = ["Sí", "No"]
POR = ["par", "docena"]
TIPOS_PAGO_TXT = {"destajo": "destajo", "fijo": "fijo", "mixto": "mixto"}

# Cada hoja: clave, título de la pestaña, descripción, columnas.
# Columna: (encabezado, clave, obligatoria, tipo, ayuda, lista)
HOJAS = [
    ("materiales", "Materiales", "Cueros, forros, suelas, pegantes, hilos, empaque… con lo que hay hoy en bodega.", [
        ("Nombre", "nombre", True, "texto", "Ej: Cuero napa negro", None),
        ("Categoría", "categoria", False, "lista", "Elige de la lista", CATEGORIAS_INSUMO),
        ("Unidad", "unidad", True, "lista", "metro cuadrado, par, kilogramo…", UNIDADES),
        ("Costo por unidad", "costo", False, "numero", "Pesos, sin puntos. Ej: 85000", None),
        ("Cantidad en bodega", "cantidad", False, "numero", "Lo que hay hoy. Ej: 12.5", None),
        ("Stock mínimo", "minimo", False, "numero", "Para la alerta de agotado", None),
    ]),
    ("referencias", "Referencias", "El catálogo: cada modelo y color es una fila.", [
        ("Código", "codigo", True, "texto", "Ej: REF-101", None),
        ("Modelo", "nombre", True, "texto", "Ej: Baleta clásica", None),
        ("Color", "color", False, "texto", "Ej: Negro", None),
        ("Línea", "linea", False, "lista", "Dama, Caballero…", LINEAS),
        ("Material", "material", False, "texto", "Ej: Cuero napa", None),
        ("Talla mínima", "tmin", True, "entero", "Ej: 35", None),
        ("Talla máxima", "tmax", True, "entero", "Ej: 40", None),
        ("Precio al detal", "precio", True, "numero", "Por par. Ej: 95000", None),
        ("Precio por mayor", "precio_mayor", False, "numero", "Por par. Ej: 68000", None),
        ("Stock mínimo (pares)", "minimo", False, "entero", "Opcional", None),
    ]),
    ("ficha", "Ficha de consumo", "Cuánto material lleva UN par de cada referencia.", [
        ("Código", "codigo", True, "texto", "Código de la referencia", None),
        ("Color", "color", False, "texto", "Si el código se repite en varios colores", None),
        ("Material", "material", True, "texto", "Igual que en la hoja Materiales", None),
        ("Cantidad por par", "cantidad", True, "numero", "Ej: 0.12", None),
        ("Unidad", "unidad", False, "lista", "Vacío = la del material", UNIDADES),
    ]),
    ("mano_obra", "Mano de obra", "Lo que se paga por cada proceso de cada referencia.", [
        ("Código", "codigo", True, "texto", "Código de la referencia", None),
        ("Color", "color", False, "texto", "Opcional", None),
        ("Proceso", "proceso", True, "lista", "Corte, Montaje…", PROCESOS),
        ("Valor", "valor", True, "numero", "Pesos. Ej: 1800", None),
        ("Por", "por", False, "lista", "par o docena (vacío = par)", POR),
    ]),
    ("existencias", "Existencias calzado", "Pares terminados que hay hoy en bodega, por talla.", [
        ("Código", "codigo", True, "texto", "Código de la referencia", None),
        ("Color", "color", False, "texto", "Opcional", None),
        ("Costo por par", "costo", False, "numero", "Vacío = costo de la ficha", None),
    ] + [(f"T{t}", f"t{t}", False, "entero", "Pares", None) for t in TALLAS]),
    ("empleados", "Empleados", "Operarios y personal de la fábrica.", [
        ("Nombre", "nombre", True, "texto", "Nombre completo", None),
        ("Documento", "documento", False, "texto", "Cédula", None),
        ("Teléfono", "telefono", False, "texto", "Celular", None),
        ("Cargo", "cargo", False, "texto", "Ej: Guarnecedora", None),
        ("Proceso", "proceso", False, "lista", "El que hace normalmente", PROCESOS),
        ("Tipo de pago", "tipo_pago", False, "lista", "destajo, fijo o mixto", list(TIPOS_PAGO_TXT)),
        ("Salario mensual", "salario", False, "numero", "Si es fijo o mixto", None),
        ("Auxilio de transporte", "aux", False, "lista", "Sí / No (vacío = Sí)", SI_NO),
        ("Fecha de ingreso", "ingreso", False, "fecha", "dd/mm/aaaa", None),
    ]),
    ("clientes", "Clientes", "Almacenes, distribuidores y particulares.", [
        ("Nombre", "nombre", True, "texto", "Ej: Almacén El Paso", None),
        ("Tipo", "tipo", False, "lista", "Almacén / tienda…", TIPOS_CLIENTE),
        ("NIT o cédula", "identificacion", False, "texto", "Sin dígito de verificación", None),
        ("Teléfono", "telefono", False, "texto", "Celular para WhatsApp", None),
        ("Email", "email", False, "texto", "Opcional", None),
        ("Dirección", "direccion", False, "texto", "Opcional", None),
        ("Ciudad", "ciudad", False, "texto", "Opcional", None),
        ("Plazo de crédito (días)", "plazo", False, "entero", "0 = contado", None),
        ("Cupo de crédito", "cupo", False, "numero", "Opcional", None),
        ("Precio por mayor", "mayor", False, "lista", "Sí / No", SI_NO),
    ]),
    ("cartera", "Cartera pendiente", "Lo que los clientes deben HOY. Una fila por factura pendiente.", [
        ("Cliente", "cliente", True, "texto", "Nombre o NIT del cliente", None),
        ("Factura N°", "factura", False, "texto", "El número que usaban antes", None),
        ("Fecha de la factura", "fecha", True, "fecha", "dd/mm/aaaa", None),
        ("Valor pendiente", "valor", True, "numero", "Lo que falta por pagar", None),
        ("Vence", "vence", False, "fecha", "Vacío = fecha + plazo del cliente", None),
    ]),
    ("ventas", "Ventas históricas", "Opcional: ventas ya pagadas de meses anteriores, para los reportes.", [
        ("Fecha", "fecha", True, "fecha", "dd/mm/aaaa", None),
        ("Factura N°", "factura", False, "texto", "Las filas con el mismo número se agrupan", None),
        ("Cliente", "cliente", False, "texto", "Opcional", None),
        ("Código", "codigo", True, "texto", "Código de la referencia", None),
        ("Color", "color", False, "texto", "Opcional", None),
        ("Talla", "talla", False, "entero", "Opcional", None),
        ("Pares", "pares", True, "entero", "Ej: 6", None),
        ("Precio por par", "precio", True, "numero", "Al que se vendió", None),
        ("Costo por par", "costo", False, "numero", "Vacío = costo de la ficha", None),
    ]),
    ("gastos", "Gastos", "Opcional: gastos de meses anteriores (arriendo, servicios…).", [
        ("Fecha", "fecha", True, "fecha", "dd/mm/aaaa", None),
        ("Categoría", "categoria", False, "lista", "Elige de la lista", CATEGORIAS_GASTO),
        ("Descripción", "descripcion", False, "texto", "Ej: Arriendo bodega marzo", None),
        ("Valor", "valor", True, "numero", "Pesos", None),
    ]),
    ("maquinas", "Máquinas", "Máquinas y equipos, para calcular su desgaste mensual.", [
        ("Nombre", "nombre", True, "texto", "Ej: Máquina de poste", None),
        ("Fecha de compra", "fecha", True, "fecha", "dd/mm/aaaa", None),
        ("Valor de compra", "valor", True, "numero", "Pesos", None),
        ("Vida útil (meses)", "vida", True, "entero", "Ej: 60", None),
        ("Valor residual", "residual", False, "numero", "Opcional", None),
        ("Proveedor", "proveedor", False, "texto", "Opcional", None),
        ("Descripción", "descripcion", False, "texto", "Opcional", None),
    ]),
]
HOJA = {k: (titulo, desc, cols) for k, titulo, desc, cols in HOJAS}


# ------------------------------------------------------------------ utilidades
def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def a_numero(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("$", "").replace(" ", "").replace("COP", "")
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".") if re.fullmatch(r"-?\d+,\d{1,2}", s) else s.replace(",", "")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        raise ValueError(f"«{v}» no es un número")


def a_fecha(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:  # número de serie de Excel
        return date(1899, 12, 30) + timedelta(days=int(v))
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"«{v}» no es una fecha (usa dd/mm/aaaa)")


def txt(v):
    """Texto de una celda; los números enteros (cédulas, teléfonos) sin «.0»."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def a_bool(v, defecto):
    if v is None or str(v).strip() == "":
        return defecto
    return norm(v) in ("si", "s", "x", "1", "true", "verdadero", "yes")


def de_lista(v, opciones):
    """Devuelve la opción que coincide sin importar tildes ni mayúsculas, o None."""
    n = norm(v)
    for o in opciones:
        if norm(o) == n:
            return o
    return None


UNIDAD_ALIAS = {"m2": "metro cuadrado", "mt2": "metro cuadrado", "metros cuadrados": "metro cuadrado",
                "dm2": "decímetro cuadrado", "pie2": "pie cuadrado", "pies cuadrados": "pie cuadrado",
                "m": "metro", "mt": "metro", "mts": "metro", "metros": "metro", "cm": "centímetro",
                "kg": "kilogramo", "kilo": "kilogramo", "kilos": "kilogramo", "kilogramos": "kilogramo",
                "g": "gramo", "gr": "gramo", "gramos": "gramo", "lb": "libra", "libras": "libra",
                "l": "litro", "lt": "litro", "litros": "litro", "ml": "mililitro", "gal": "galón",
                "galon": "galón", "galones": "galón", "und": "unidad", "un": "unidad", "unidades": "unidad",
                "u": "unidad", "pares": "par", "pr": "par", "docenas": "docena", "doc": "docena",
                "rollos": "rollo", "laminas": "lámina", "conos": "cono", "paquetes": "paquete",
                "cajas": "caja", "bolsas": "bolsa", "yardas": "yarda", "yd": "yarda"}


def a_unidad(v):
    if v is None or str(v).strip() == "":
        return None
    u = de_lista(v, UNIDADES)
    return u or UNIDAD_ALIAS.get(norm(v))


# ------------------------------------------------------------------ plantilla
AZUL = "1F4E79"
CAFE = "8A4B2A"


def plantilla(ejemplo=False):
    """Excel para llenar. Con ejemplo=True trae datos de muestra de una fábrica ficticia."""
    wb = Workbook()
    ins = wb.active
    ins.title = "Instrucciones"
    listas = wb.create_sheet("Listas")
    listas.sheet_state = "hidden"
    rangos = {}
    col = 1
    for nombre, ops in (("unidades", UNIDADES), ("cat_ins", CATEGORIAS_INSUMO), ("lineas", LINEAS),
                        ("procesos", PROCESOS), ("tipos_cli", TIPOS_CLIENTE), ("cat_gasto", CATEGORIAS_GASTO),
                        ("si_no", SI_NO), ("por", POR), ("tipo_pago", list(TIPOS_PAGO_TXT))):
        for i, o in enumerate(ops, 1):
            listas.cell(row=i, column=col, value=o)
        L = get_column_letter(col)
        rangos[tuple(ops)] = f"Listas!${L}$1:${L}${len(ops)}"
        col += 1

    borde = Border(bottom=Side(style="thin", color="BBBBBB"))
    for k, titulo, desc, cols in HOJAS:
        ws = wb.create_sheet(titulo)
        ws.freeze_panes = "A3"
        for j, (enc, clave, req, tipo, ayuda, ops) in enumerate(cols, 1):
            c = ws.cell(row=1, column=j, value=enc + (" *" if req else ""))
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=CAFE if req else AZUL)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            h = ws.cell(row=2, column=j, value=f"{MARCA_AYUDA} {ayuda}")
            h.font = Font(italic=True, color="7F7F7F", size=9)
            h.fill = PatternFill("solid", fgColor="F3F3F3")
            h.alignment = Alignment(wrap_text=True, vertical="top")
            h.border = borde
            if enc.startswith("T") and enc[1:].isdigit():
                ancho = 7
            elif tipo in ("texto", "lista"):
                ancho = 30 if clave in ("nombre", "material", "descripcion", "cliente", "categoria", "tipo",
                                        "direccion", "proceso") else 20
            else:
                ancho = max(14, len(enc) + 4)
            ws.column_dimensions[get_column_letter(j)].width = ancho
            if ops:
                dv = DataValidation(type="list", formula1=rangos[tuple(ops)], allow_blank=True,
                                    showErrorMessage=False)
                dv.add(f"{get_column_letter(j)}3:{get_column_letter(j)}{MAX_FILAS}")
                ws.add_data_validation(dv)
        ws.row_dimensions[1].height = 32
        ws.row_dimensions[2].height = 30
        ws.cell(row=1, column=1).comment = Comment(desc, "SoftGest")
        if ejemplo:
            for r, fila in enumerate(EJEMPLO.get(k, []), 3):
                for j, v in enumerate(fila, 1):
                    if v is not None:
                        cel = ws.cell(row=r, column=j, value=v)
                        if isinstance(v, date):
                            cel.number_format = "DD/MM/YYYY"

    # Hoja de instrucciones
    ins.column_dimensions["A"].width = 4
    ins.column_dimensions["B"].width = 30
    ins.column_dimensions["C"].width = 90
    ins["B2"] = "Plantilla de importación · CalzadoControl"
    ins["B2"].font = Font(bold=True, size=16, color=CAFE)
    pasos = [
        "1. Llena solo las hojas que necesites. Las que dejes vacías no se tocan.",
        "2. Copia y pega desde tus Excel debajo de la fila gris de ayuda (desde la fila 3). No cambies los encabezados.",
        "3. Las columnas en café con * son obligatorias. Las azules son opcionales.",
        "4. Números sin signos: 85000 o 85.000 sirven. Fechas como 15/03/2026.",
        "5. Las referencias se reconocen por Código (+ Color si el mismo código tiene varios colores). "
        "Los materiales por su nombre. Los clientes por NIT o nombre. Los empleados por documento o nombre.",
        "6. Si algo ya existe en la app, se ACTUALIZA (no se duplica). Las cantidades de bodega se SUMAN.",
        "7. En la app: Datos del negocio → Importar datos → sube el archivo. Verás una revisión con los "
        "errores fila por fila antes de guardar. Si hay errores no se guarda nada.",
        "8. Si te equivocas, en la misma pantalla puedes Deshacer la importación completa.",
    ]
    for i, p in enumerate(pasos, 4):
        ins.cell(row=i, column=2, value=p).alignment = Alignment(wrap_text=True)
        ins.merge_cells(start_row=i, start_column=2, end_row=i, end_column=3)
        ins.row_dimensions[i].height = 30
    r = 4 + len(pasos) + 1
    ins.cell(row=r, column=2, value="Hoja").font = Font(bold=True)
    ins.cell(row=r, column=3, value="Para qué sirve").font = Font(bold=True)
    for k, titulo, desc, cols in HOJAS:
        r += 1
        ins.cell(row=r, column=2, value=titulo).font = Font(bold=True, color=AZUL)
        ins.cell(row=r, column=3, value=desc).alignment = Alignment(wrap_text=True)
    r += 2
    ins.cell(row=r, column=2, value="Importante").font = Font(bold=True, color="C00000")
    for t in ("Existencias calzado y Cantidad en bodega de Materiales: escribe lo que hay HOY. "
              "No restan materiales ni cuentan como producción del mes.",
              "Cartera pendiente: solo lo que te deben hoy. Aparece en Cartera para cobrar y registrar abonos; "
              "no se suma otra vez a las ventas.",
              "Ventas históricas: no mueven el inventario (para eso está la hoja Existencias). "
              "Solo alimentan los reportes de meses anteriores."):
        r += 1
        c = ins.cell(row=r, column=2, value="• " + t)
        c.alignment = Alignment(wrap_text=True)
        ins.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ins.row_dimensions[r].height = 32
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    ins.sheet_view.showGridLines = False
    wb.move_sheet("Listas", offset=len(wb.sheetnames))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_h = date.today()
EJEMPLO = {
    "materiales": [
        ["Cuero napa negro", "Cueros y sintéticos", "metro cuadrado", 92000, 18.5, 5],
        ["Cuero napa café", "Cueros y sintéticos", "metro cuadrado", 92000, 11, 5],
        ["Forro badana", "Forros", "metro cuadrado", 38000, 14, 4],
        ["Suela TR dama 35-40", "Suelas y tacones", "par", 6500, 240, 60],
        ["Pegante PU", "Pegantes y químicos", "galón", 95000, 3, 1],
        ["Hilo nylon 40", "Hilos", "cono", 14000, 9, 3],
        ["Caja de cartón", "Empaque", "unidad", 1200, 400, 100],
    ],
    "referencias": [
        ["REF-101", "Baleta clásica", "Negro", "Dama", "Cuero napa", 35, 40, 95000, 68000, 12],
        ["REF-101", "Baleta clásica", "Café", "Dama", "Cuero napa", 35, 40, 95000, 68000, 12],
        ["REF-205", "Botín casual", "Negro", "Caballero", "Cuero napa", 38, 43, 165000, 115000, 6],
    ],
    "ficha": [
        ["REF-101", "Negro", "Cuero napa negro", 0.12, "metro cuadrado"],
        ["REF-101", "Negro", "Forro badana", 8, "decímetro cuadrado"],
        ["REF-101", "Negro", "Suela TR dama 35-40", 1, None],
        ["REF-101", "Negro", "Pegante PU", 0.004, None],
        ["REF-101", "Negro", "Caja de cartón", 1, None],
        ["REF-101", "Café", "Cuero napa café", 0.12, "metro cuadrado"],
        ["REF-101", "Café", "Suela TR dama 35-40", 1, None],
    ],
    "mano_obra": [
        ["REF-101", "Negro", "Corte", 1800, "par"], ["REF-101", "Negro", "Guarnición (aparado)", 30000, "docena"],
        ["REF-101", "Negro", "Montaje", 2200, "par"], ["REF-101", "Café", "Corte", 1800, "par"],
        ["REF-205", "Negro", "Montaje", 3500, "par"],
    ],
    "existencias": [
        ["REF-101", "Negro", 41000] + [None] * 14 + [6, 12, 18, 15, 9, 4] + [None] * 5,
        ["REF-101", "Café", None] + [None] * 14 + [3, 8, 10, 9, 5, 2] + [None] * 5,
        ["REF-205", "Negro", 60000] + [None] * 17 + [4, 7, 9, 8, 5, 2] + [None] * 2,
    ],
    "empleados": [
        ["Luz Marina Peña", "60321456", "3104567890", "Guarnecedora", "Guarnición (aparado)", "destajo", None, "Sí",
         date(2023, 2, 1)],
        ["Carlos Rincón", "88234567", "3157894561", "Montador", "Montaje", "mixto", 1750905, "Sí", date(2021, 8, 16)],
        ["Andrea Suárez", "1090456789", "3201234567", "Administradora", None, "fijo", 2200000, "No", date(2022, 1, 10)],
    ],
    "clientes": [
        ["Almacén El Paso", "Almacén / tienda", "900123456", "3001234567", None, "Cl 10 # 5-20", "Cúcuta", 30,
         3000000, "Sí"],
        ["Calzado La 13", "Distribuidor", "901555777", "3112223344", "compras@la13.co", None, "Bucaramanga", 45,
         8000000, "Sí"],
        ["Marta Gómez", "Particular", "37555111", "3209998877", None, None, "Cúcuta", 0, None, "No"],
    ],
    "cartera": [
        ["Almacén El Paso", "F-1520", _h - timedelta(days=40), 1360000, None],
        ["Calzado La 13", "F-1544", _h - timedelta(days=12), 2450000, _h + timedelta(days=33)],
    ],
    "ventas": [
        [_h - timedelta(days=70), "F-1401", "Almacén El Paso", "REF-101", "Negro", 37, 6, 68000, 40500],
        [_h - timedelta(days=70), "F-1401", "Almacén El Paso", "REF-101", "Negro", 38, 6, 68000, 40500],
        [_h - timedelta(days=45), "F-1478", "Marta Gómez", "REF-205", "Negro", 41, 1, 165000, None],
    ],
    "gastos": [
        [_h - timedelta(days=60), "Arriendo", "Arriendo bodega", 1800000],
        [_h - timedelta(days=55), "Energía eléctrica", "Factura Centrales Eléctricas", 420000],
        [_h - timedelta(days=30), "Arriendo", "Arriendo bodega", 1800000],
    ],
    "maquinas": [
        ["Máquina de poste", date(2022, 3, 1), 4800000, 84, 400000, "Maquinaria del Norte", None],
        ["Troqueladora", date(2020, 6, 15), 9500000, 120, 1000000, None, "Hidráulica 20 t"],
    ],
}


# ------------------------------------------------------------------ lectura
def leer(contenido):
    """Lee el Excel y devuelve {clave_hoja: [(n_fila, {clave: valor})]}, y avisos."""
    try:
        wb = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    except Exception:
        raise ErrorNegocio("No se pudo leer el archivo. Debe ser un Excel .xlsx (guárdalo como «Libro de Excel»).")
    por_titulo = {norm(ws.title): ws for ws in wb.worksheets}
    datos, avisos = {}, []
    for k, titulo, desc, cols in HOJAS:
        ws = por_titulo.get(norm(titulo))
        if ws is None:
            continue
        filas = ws.iter_rows(values_only=True)
        enc = next(filas, None)
        if not enc:
            continue
        mapa = {}
        for j, h in enumerate(enc):
            hn = norm(str(h or "").replace("*", ""))
            for e, clave, *_ in cols:
                if hn and hn == norm(e):
                    mapa[j] = clave
        faltan = [e for e, clave, req, *_ in cols if req and clave not in mapa.values()]
        if faltan:
            avisos.append(f"Hoja «{ws.title}»: faltan las columnas {', '.join(faltan)}. Se ignoró la hoja.")
            continue
        lista = []
        for n, fila in enumerate(filas, 2):
            if n > MAX_FILAS:
                avisos.append(f"Hoja «{ws.title}»: solo se leen las primeras {MAX_FILAS} filas.")
                break
            if any(isinstance(v, str) and v.strip().startswith(MARCA_AYUDA) for v in fila):
                continue
            reg = {mapa[j]: (v.strip() if isinstance(v, str) else v)
                   for j, v in enumerate(fila) if j in mapa}
            if all(v is None or v == "" for v in reg.values()):
                continue
            lista.append((n, reg))
        if lista:
            datos[k] = lista
    wb.close()
    return datos, avisos


# ------------------------------------------------------------------ aplicación
class Importacion:
    """Aplica los datos dentro de la sesión actual. No hace commit."""

    def __init__(self, datos, usuario_id, fecha_saldos):
        self.datos = datos
        self.usuario_id = usuario_id
        self.fecha = fecha_saldos or ahora().date()
        self.errores = []
        self.creado = []  # [tipo, id, ...] para deshacer
        self.res = {}  # clave_hoja -> {nuevos, actualizados, ...}
        self._insumos = {norm(i.nombre): i for i in Insumo.query.all()}
        self._prods = Producto.query.all()
        self._clientes = Cliente.query.all()
        self._empleados = Empleado.query.all()
        self._pendientes = []  # (tipo, objeto, extra) cuyo id se conoce al hacer flush
        self._ref_mala = {}  # (código, color) -> fila con error en la hoja Referencias
        self._mat_mala = {}  # nombre -> fila con error en la hoja Materiales

    # --- helpers
    def err(self, hoja, fila, msg):
        self.errores.append((HOJA[hoja][0], fila, msg))

    def cuenta(self, hoja, que, n=1):
        d = self.res.setdefault(hoja, {})
        d[que] = d.get(que, 0) + n

    def anotar(self, tipo, obj, *extra):
        self._pendientes.append((tipo, obj, extra))

    def val(self, hoja, fila, reg, clave, tipo, req=False):
        v = reg.get(clave)
        vacio = v is None or str(v).strip() == ""
        titulo_col = next(e for e, c, *_ in HOJA[hoja][2] if c == clave)
        if vacio:
            if req:
                self.err(hoja, fila, f"falta «{titulo_col}».")
                raise _FilaMala
            return None
        try:
            if tipo == "numero":
                x = a_numero(v)
                if x < 0:
                    raise ValueError(f"«{titulo_col}» no puede ser negativo")
                return x
            if tipo == "entero":
                x = a_numero(v)
                if x != int(x) or x < 0:
                    raise ValueError(f"«{titulo_col}» debe ser un número entero")
                return int(x)
            if tipo == "fecha":
                return a_fecha(v)
            return txt(v)
        except ValueError as e:
            self.err(hoja, fila, f"{titulo_col}: {e}.")
            raise _FilaMala

    def producto(self, hoja, fila, codigo, color):
        c = norm(codigo)
        cands = [p for p in self._prods if norm(p.codigo) == c]
        if color:
            cands = [p for p in cands if norm(p.color) == norm(color)]
        if not cands:
            mala = self._ref_mala.get((c, norm(color))) or next(
                (f for (cc, _), f in self._ref_mala.items() if cc == c and not color), None)
            if mala:
                self.err(hoja, fila, f"la referencia «{codigo}{' ' + color if color else ''}» tiene un error en "
                                     f"la hoja Referencias (fila {mala}); corrígelo allá primero.")
                raise _FilaMala
            self.err(hoja, fila, f"no existe la referencia «{codigo}{' ' + color if color else ''}». "
                                 "Créala en la hoja Referencias.")
            raise _FilaMala
        if len(cands) > 1:
            self.err(hoja, fila, f"el código «{codigo}» tiene varios colores; escribe el Color.")
            raise _FilaMala
        return cands[0]

    def cliente(self, texto, crear_hoja=None):
        n = norm(texto)
        for c in self._clientes:
            if n and (norm(c.nombre) == n or (c.identificacion and norm(c.identificacion) == n)):
                return c
        if crear_hoja:
            c = Cliente(nombre=str(texto)[:120])
            db.session.add(c)
            self._clientes.append(c)
            self.anotar("cliente_nuevo", c)
            self.cuenta(crear_hoja, "clientes_creados")
            return c
        return None

    # --- secciones
    def materiales(self):
        h = "materiales"
        for fila, reg in self.datos.get(h, []):
            antes = len(self.errores)
            self._materiales_fila(h, fila, reg)
            if len(self.errores) > antes and reg.get("nombre"):
                self._mat_mala[norm(reg.get("nombre"))] = fila

    def _materiales_fila(self, h, fila, reg):
        if True:
            try:
                nombre = self.val(h, fila, reg, "nombre", "texto", True)
                unidad = a_unidad(reg.get("unidad"))
                if not unidad:
                    self.err(h, fila, f"la unidad «{reg.get('unidad') or ''}» no existe. Usa una de la lista "
                                      "(metro cuadrado, par, kilogramo, litro, unidad…).")
                    return
                cat = de_lista(reg.get("categoria"), CATEGORIAS_INSUMO) if reg.get("categoria") else None
                if reg.get("categoria") and not cat:
                    cat = "Otros"
                costo = self.val(h, fila, reg, "costo", "numero")
                cant = self.val(h, fila, reg, "cantidad", "numero")
                minimo = self.val(h, fila, reg, "minimo", "numero")
                ins = self._insumos.get(norm(nombre))
                if ins and ins.es_preparacion:
                    self.err(h, fila, f"«{nombre}» es un semielaborado; su existencia se maneja en Semielaborados.")
                    return
                if ins is None:
                    ins = Insumo(nombre=nombre[:120], unidad=unidad, categoria=cat or "Otros",
                                 stock=0, costo_promedio=costo or 0, stock_minimo=minimo or 0)
                    db.session.add(ins)
                    self._insumos[norm(nombre)] = ins
                    self.anotar("insumo_nuevo", ins)
                    self.cuenta(h, "nuevos")
                else:
                    if ins.unidad != unidad and (ins.stock or 0) > 0:
                        self.err(h, fila, f"«{nombre}» ya existe en {ins.unidad} con existencias; "
                                          "no se le puede cambiar la unidad.")
                        return
                    ins.unidad = unidad
                    if cat:
                        ins.categoria = cat
                    if minimo is not None:
                        ins.stock_minimo = minimo
                    if costo and not cant and (ins.stock or 0) <= 0:
                        ins.costo_promedio = costo
                    ins.activo = True
                    self.cuenta(h, "actualizados")
                if cant:
                    if not costo and not ins.costo_promedio:
                        self.err(h, fila, f"escribe el «Costo por unidad» de «{nombre}» para valorizar lo que hay en bodega.")
                        return
                    previo = ins.costo_promedio or 0
                    ins.costo_promedio = promedio_ponderado(ins.stock, previo, cant, costo or previo)
                    ins.stock = (ins.stock or 0) + cant
                    self.anotar("insumo_stock", ins, cant, previo)
                    self.cuenta(h, "valor_bodega", cant * (costo or previo))
            except _FilaMala:
                pass

    def referencias(self):
        h = "referencias"
        for fila, reg in self.datos.get(h, []):
            antes = len(self.errores)
            self._referencias_fila(h, fila, reg)
            if len(self.errores) > antes and reg.get("codigo"):
                self._ref_mala[(norm(reg.get("codigo")), norm(reg.get("color")))] = fila

    def _referencias_fila(self, h, fila, reg):
        if True:
            try:
                codigo = self.val(h, fila, reg, "codigo", "texto", True)
                nombre = self.val(h, fila, reg, "nombre", "texto", True)
                color = self.val(h, fila, reg, "color", "texto") or ""
                tmin = self.val(h, fila, reg, "tmin", "entero", True)
                tmax = self.val(h, fila, reg, "tmax", "entero", True)
                precio = self.val(h, fila, reg, "precio", "numero", True)
                mayor = self.val(h, fila, reg, "precio_mayor", "numero")
                minimo = self.val(h, fila, reg, "minimo", "entero")
                if not (15 <= tmin <= tmax <= 50):
                    self.err(h, fila, f"rango de tallas {tmin}–{tmax} no válido.")
                    return
                linea = de_lista(reg.get("linea"), LINEAS) or ("Dama" if not reg.get("linea") else "Unisex")
                p = next((x for x in self._prods if norm(x.codigo) == norm(codigo) and norm(x.color) == norm(color)), None)
                if p is None:
                    p = Producto(codigo=codigo.upper()[:30], nombre=nombre[:120], color=color[:60] or None,
                                 linea=linea, material=(reg.get("material") or None), talla_min=tmin,
                                 talla_max=tmax, precio_venta=precio, precio_mayor=mayor or 0,
                                 stock_minimo=minimo or 0, stock=0, costo_promedio=0, tanda_rinde=1)
                    db.session.add(p)
                    self._prods.append(p)
                    self.anotar("producto_nuevo", p)
                    self.cuenta(h, "nuevos")
                else:
                    con_stock = [t for t, s in p.stock_por_talla.items() if s and not (tmin <= t <= tmax)]
                    if con_stock:
                        self.err(h, fila, f"«{p.etiqueta}» tiene pares en tallas {con_stock} fuera del rango {tmin}–{tmax}.")
                        return
                    p.nombre, p.talla_min, p.talla_max, p.precio_venta = nombre[:120], tmin, tmax, precio
                    if mayor is not None:
                        p.precio_mayor = mayor
                    if reg.get("linea"):
                        p.linea = linea
                    if reg.get("material"):
                        p.material = str(reg["material"])[:80]
                    if minimo is not None:
                        p.stock_minimo = minimo
                    p.activo = True
                    self.cuenta(h, "actualizados")
            except _FilaMala:
                pass

    def ficha(self):
        h = "ficha"
        nuevas = {}  # producto -> [(insumo, cant, unidad)]
        for fila, reg in self.datos.get(h, []):
            try:
                p = self.producto(h, fila, self.val(h, fila, reg, "codigo", "texto", True), reg.get("color"))
                mat = self.val(h, fila, reg, "material", "texto", True)
                cant = self.val(h, fila, reg, "cantidad", "numero", True)
                ins = self._insumos.get(norm(mat))
                if ins is None and norm(mat) in self._mat_mala:
                    self.err(h, fila, f"el material «{mat}» tiene un error en la hoja Materiales "
                                      f"(fila {self._mat_mala[norm(mat)]}); corrígelo allá primero.")
                    continue
                if ins is None:
                    self.err(h, fila, f"no existe el material «{mat}». Agrégalo en la hoja Materiales.")
                    continue
                u = a_unidad(reg.get("unidad")) or ins.unidad
                if u != ins.unidad:
                    fu, fi = FACTORES.get(u), FACTORES.get(ins.unidad)
                    if not fu or not fi or fu[0] != fi[0]:
                        self.err(h, fila, f"«{mat}» se maneja en {ins.unidad}; no se puede usar {u}.")
                        continue
                if cant <= 0:
                    self.err(h, fila, "la cantidad por par debe ser mayor a 0.")
                    continue
                nuevas.setdefault(p, []).append((ins, convertir(cant, u, ins.unidad), u))
            except _FilaMala:
                pass
        for p, lineas in nuevas.items():
            for r in list(p.receta):
                p.receta.remove(r)
            for ins, cant, u in lineas:
                r = RecetaItem(insumo=ins, cantidad=cant, unidad=u)
                p.receta.append(r)
                self.anotar("receta", r)
            p.tanda_rinde = 1
            self.cuenta(h, "referencias")
            self.cuenta(h, "lineas", len(lineas))

    def mano_obra(self):
        h = "mano_obra"
        nuevas = {}
        for fila, reg in self.datos.get(h, []):
            try:
                p = self.producto(h, fila, self.val(h, fila, reg, "codigo", "texto", True), reg.get("color"))
                proc = de_lista(reg.get("proceso"), PROCESOS)
                if not proc:
                    self.err(h, fila, f"el proceso «{reg.get('proceso') or ''}» no existe. Usa: {', '.join(PROCESOS)}.")
                    continue
                valor = self.val(h, fila, reg, "valor", "numero", True)
                por = de_lista(reg.get("por"), POR) or "par"
                nuevas.setdefault(p, {})[proc] = valor / 12 if por == "docena" else valor
            except _FilaMala:
                pass
        for p, tarifas in nuevas.items():
            for t in list(p.procesos):
                if t.proceso in tarifas:
                    p.procesos.remove(t)
            for proc, v in tarifas.items():
                t = ProcesoTarifa(proceso=proc, valor_par=v)
                p.procesos.append(t)
                self.anotar("tarifa", t)
            self.cuenta(h, "tarifas", len(tarifas))

    def existencias(self):
        h = "existencias"
        for fila, reg in self.datos.get(h, []):
            try:
                p = self.producto(h, fila, self.val(h, fila, reg, "codigo", "texto", True), reg.get("color"))
                corrida = {}
                for t in TALLAS:
                    c = self.val(h, fila, reg, f"t{t}", "entero")
                    if c:
                        if not (p.talla_min <= t <= p.talla_max):
                            self.err(h, fila, f"«{p.etiqueta}» va de la talla {p.talla_min} a la {p.talla_max}; "
                                              f"tiene pares en la T{t}.")
                            raise _FilaMala
                        corrida[t] = c
                if not corrida:
                    continue
                costo = self.val(h, fila, reg, "costo", "numero") or p.costo_estandar or p.costo_promedio
                if not costo:
                    self.err(h, fila, f"escribe el «Costo por par» de «{p.etiqueta}» (no tiene ficha de consumo para calcularlo).")
                    continue
                q = sum(corrida.values())
                prod = registrar_produccion(p, corrida, {}, [], otros=costo * q, otros_descripcion="Inventario inicial",
                                            nota="Inventario inicial (importado de Excel)",
                                            fecha=datetime.combine(self.fecha, datetime.min.time()).replace(hour=8),
                                            usuario_id=self.usuario_id, sin_materiales=True, es_inicial=True)
                self.anotar("produccion", prod)
                self.cuenta(h, "pares", q)
                self.cuenta(h, "valor", costo * q)
            except _FilaMala:
                pass
            except ErrorNegocio as e:
                for m in e.errores:
                    self.err(h, fila, m)

    def empleados(self):
        h = "empleados"
        for fila, reg in self.datos.get(h, []):
            try:
                nombre = self.val(h, fila, reg, "nombre", "texto", True)
                doc = txt(reg.get("documento"))[:40] or None
                tipo = de_lista(reg.get("tipo_pago"), list(TIPOS_PAGO_TXT)) or "destajo"
                salario = self.val(h, fila, reg, "salario", "numero") or 0
                if tipo in ("fijo", "mixto") and not salario:
                    self.err(h, fila, f"«{nombre}» es de pago {tipo}: escribe el salario mensual.")
                    continue
                ingreso = self.val(h, fila, reg, "ingreso", "fecha")
                e = next((x for x in self._empleados if (doc and x.documento and norm(x.documento) == norm(doc))
                          or (not doc and norm(x.nombre) == norm(nombre))), None)
                nuevo = e is None
                if nuevo:
                    e = Empleado(nombre=nombre[:120])
                    db.session.add(e)
                    self._empleados.append(e)
                    self.anotar("empleado_nuevo", e)
                e.nombre, e.tipo_pago, e.salario = nombre[:120], tipo, salario
                e.documento = doc or e.documento
                e.telefono = txt(reg.get("telefono"))[:40] or e.telefono
                e.cargo = str(reg.get("cargo") or e.cargo or "")[:80] or None
                e.proceso = de_lista(reg.get("proceso"), PROCESOS) or e.proceso
                e.aux_transporte = a_bool(reg.get("aux"), True)
                e.fecha_ingreso = ingreso or e.fecha_ingreso
                e.activo = True
                self.cuenta(h, "nuevos" if nuevo else "actualizados")
            except _FilaMala:
                pass

    def clientes(self):
        h = "clientes"
        for fila, reg in self.datos.get(h, []):
            try:
                nombre = self.val(h, fila, reg, "nombre", "texto", True)
                ident = txt(reg.get("identificacion"))[:40] or None
                plazo = self.val(h, fila, reg, "plazo", "entero")
                cupo = self.val(h, fila, reg, "cupo", "numero")
                c = next((x for x in self._clientes if (ident and x.identificacion and norm(x.identificacion) == norm(ident))
                          or norm(x.nombre) == norm(nombre)), None)
                nuevo = c is None
                if nuevo:
                    c = Cliente(nombre=nombre[:120])
                    db.session.add(c)
                    self._clientes.append(c)
                    self.anotar("cliente_nuevo", c)
                c.nombre = nombre[:120]
                c.tipo = de_lista(reg.get("tipo"), TIPOS_CLIENTE) or c.tipo or "Almacén / tienda"
                c.identificacion = ident or c.identificacion
                for k, largo in (("telefono", 40), ("email", 120), ("direccion", 160), ("ciudad", 80)):
                    if reg.get(k):
                        setattr(c, k, txt(reg[k])[:largo])
                if plazo is not None:
                    c.plazo_dias = plazo
                if cupo is not None:
                    c.cupo = cupo
                if reg.get("mayor") not in (None, ""):
                    c.usa_precio_mayor = a_bool(reg.get("mayor"), False)
                c.activo = True
                self.cuenta(h, "nuevos" if nuevo else "actualizados")
            except _FilaMala:
                pass

    def cartera(self):
        h = "cartera"
        for fila, reg in self.datos.get(h, []):
            try:
                quien = self.val(h, fila, reg, "cliente", "texto", True)
                fecha = self.val(h, fila, reg, "fecha", "fecha", True)
                valor = self.val(h, fila, reg, "valor", "numero", True)
                vence = self.val(h, fila, reg, "vence", "fecha")
                if valor <= 0:
                    continue
                c = self.cliente(quien, crear_hoja=h)
                vence = vence or fecha + timedelta(days=c.plazo_dias or 30)
                fac = self.val(h, fila, reg, "factura", "texto")
                if c.id and Venta.query.filter(Venta.origen == "saldo_inicial", Venta.anulada.is_(False),
                                               Venta.cliente_id == c.id, Venta.total == valor,
                                               db.func.date(Venta.fecha) == fecha).count():
                    self.cuenta(h, "repetidos")
                    continue
                v = Venta(fecha=datetime.combine(fecha, datetime.min.time()).replace(hour=12), cliente=c.nombre,
                          cliente_id=c.id, cliente_doc=c.identificacion, condicion="credito", vence=vence,
                          metodo_pago="Crédito", total=valor, costo_total=0, ganancia=0, iva=0, iva_pct=0,
                          origen="saldo_inicial", usuario_id=self.usuario_id,
                          nota=f"Saldo anterior{' · factura ' + fac if fac else ''} (importado de Excel)")
                v.cliente_rel_tmp = c
                db.session.add(v)
                self.anotar("venta", v)
                self.cuenta(h, "facturas")
                self.cuenta(h, "valor", valor)
            except _FilaMala:
                pass

    def ventas(self):
        h = "ventas"
        grupos = {}
        for fila, reg in self.datos.get(h, []):
            try:
                fecha = self.val(h, fila, reg, "fecha", "fecha", True)
                fac = self.val(h, fila, reg, "factura", "texto")
                quien = self.val(h, fila, reg, "cliente", "texto")
                p = self.producto(h, fila, self.val(h, fila, reg, "codigo", "texto", True), reg.get("color"))
                talla = self.val(h, fila, reg, "talla", "entero")
                pares = self.val(h, fila, reg, "pares", "entero", True)
                precio = self.val(h, fila, reg, "precio", "numero", True)
                costo = self.val(h, fila, reg, "costo", "numero")
                if costo is None:
                    costo = p.costo_ref or 0
                if pares <= 0:
                    continue
                if fecha > ahora().date():
                    self.err(h, fila, "la fecha es futura.")
                    continue
                clave = (fac or "", fecha if not fac else None, norm(quien) if not fac else "")
                g = grupos.setdefault(clave, dict(fecha=fecha, fac=fac, cliente=quien, items=[]))
                g["items"].append((p, talla, pares, precio, costo))
            except _FilaMala:
                pass
        for g in grupos.values():
            tot_g = sum(pr * q for _, _, q, pr, _ in g["items"])
            if Venta.query.filter(Venta.origen == "historica", Venta.anulada.is_(False), Venta.total == tot_g,
                                  db.func.date(Venta.fecha) == g["fecha"]).count():
                self.cuenta(h, "repetidos")
                continue
            c = self.cliente(g["cliente"]) if g["cliente"] else None
            v = Venta(fecha=datetime.combine(g["fecha"], datetime.min.time()).replace(hour=12),
                      cliente=(c.nombre if c else g["cliente"]) or None, cliente_id=c.id if c else None,
                      condicion="contado", metodo_pago="Efectivo", iva=0, iva_pct=0, origen="historica",
                      usuario_id=self.usuario_id,
                      nota=f"Venta histórica{' · factura ' + g['fac'] if g['fac'] else ''} (importada de Excel)")
            if c is not None and c.id is None:
                v.cliente_rel_tmp = c
            tot = cos = 0
            for p, talla, pares, precio, costo in g["items"]:
                v.items.append(VentaItem(producto=p, talla=talla, cantidad=pares, precio_unitario=precio,
                                         costo_unitario=costo, subtotal=precio * pares, costo=costo * pares))
                tot += precio * pares
                cos += costo * pares
            v.total, v.costo_total, v.ganancia = tot, cos, tot - cos
            db.session.add(v)
            self.anotar("venta", v)
            self.cuenta(h, "ventas")
            self.cuenta(h, "valor", tot)

    def gastos(self):
        h = "gastos"
        for fila, reg in self.datos.get(h, []):
            try:
                fecha = self.val(h, fila, reg, "fecha", "fecha", True)
                valor = self.val(h, fila, reg, "valor", "numero", True)
                if not valor:
                    continue
                cat = de_lista(reg.get("categoria"), CATEGORIAS_GASTO) or "Otros"
                if Gasto.query.filter(Gasto.fecha == fecha, Gasto.valor == valor, Gasto.categoria == cat).count():
                    self.cuenta(h, "repetidos")
                    continue
                g = Gasto(fecha=fecha, categoria=de_lista(reg.get("categoria"), CATEGORIAS_GASTO) or "Otros",
                          descripcion=(str(reg.get("descripcion") or "") or None), valor=valor,
                          usuario_id=self.usuario_id)
                db.session.add(g)
                self.anotar("gasto", g)
                self.cuenta(h, "gastos")
                self.cuenta(h, "valor", valor)
            except _FilaMala:
                pass

    def maquinas(self):
        h = "maquinas"
        for fila, reg in self.datos.get(h, []):
            try:
                nombre = self.val(h, fila, reg, "nombre", "texto", True)
                fecha = self.val(h, fila, reg, "fecha", "fecha", True)
                valor = self.val(h, fila, reg, "valor", "numero", True)
                vida = self.val(h, fila, reg, "vida", "entero", True)
                residual = self.val(h, fila, reg, "residual", "numero") or 0
                if vida <= 0:
                    self.err(h, fila, "la vida útil debe ser mayor a 0 meses.")
                    continue
                if residual > valor:
                    self.err(h, fila, "el valor residual no puede ser mayor al valor de compra.")
                    continue
                if Activo.query.filter(Activo.nombre == nombre[:120], Activo.fecha_compra == fecha,
                                       Activo.eliminado.is_(False)).count():
                    self.cuenta(h, "repetidos")
                    continue
                a = Activo(nombre=nombre[:120], fecha_compra=fecha, valor=valor, vida_util_meses=vida,
                           valor_residual=residual, proveedor=(reg.get("proveedor") or None),
                           descripcion=(reg.get("descripcion") or None))
                db.session.add(a)
                self.anotar("activo", a)
                self.cuenta(h, "maquinas")
            except _FilaMala:
                pass

    def correr(self):
        orden = ["materiales", "referencias", "ficha", "mano_obra", "empleados", "clientes",
                 "existencias", "cartera", "ventas", "gastos", "maquinas"]
        for k in orden:
            if k in self.datos:
                if k == "existencias":
                    db.session.flush()  # costo estándar necesita la ficha guardada
                getattr(self, k)()
        db.session.flush()
        for v in [o for t, o, _ in self._pendientes if t == "venta"]:
            c = getattr(v, "cliente_rel_tmp", None)
            if c is not None:
                v.cliente_id, v.cliente_doc = c.id, c.identificacion
        db.session.flush()
        self.creado = [[t, o.id, *extra] for t, o, extra in self._pendientes]
        return not self.errores


class _FilaMala(Exception):
    pass


# ------------------------------------------------------------------ deshacer
def deshacer(items):
    """Revierte lo que creó un lote. Lanza ErrorNegocio si algo ya se usó y no se puede revertir."""
    modelos = {"insumo_nuevo": Insumo, "insumo_stock": Insumo, "producto_nuevo": Producto, "receta": RecetaItem,
               "tarifa": ProcesoTarifa, "produccion": Produccion, "empleado_nuevo": Empleado,
               "cliente_nuevo": Cliente, "venta": Venta, "gasto": Gasto, "activo": Activo}
    errores = []
    nuevos = {t: {i[1] for i in items if i[0] == t} for t in ("producto_nuevo", "insumo_nuevo", "cliente_nuevo",
                                                              "empleado_nuevo", "venta", "produccion")}
    desactivados = 0
    for it in reversed(items):
        tipo, oid, *extra = it
        obj = db.session.get(modelos[tipo], oid)
        if obj is None:
            continue
        if tipo == "venta":
            if obj.anulada:
                continue
            if any(not a.anulado for a in obj.abonos):
                errores.append(f"La factura importada de «{obj.cliente}» ya tiene abonos registrados. "
                               "Anula esos abonos en Cartera antes de deshacer.")
                continue
            db.session.delete(obj)
        elif tipo == "produccion":
            if not obj.anulada:
                try:
                    anular_produccion(obj)
                except ErrorNegocio as e:
                    errores += [f"Existencias de «{obj.producto.etiqueta}»: {m}" for m in e.errores]
        elif tipo == "insumo_stock":
            cant, previo = extra
            if (obj.stock or 0) + 1e-9 < cant:
                errores.append(f"De «{obj.nombre}» ya se usó parte de lo importado (quedan {obj.stock:g}, "
                               f"se importaron {cant:g}).")
                continue
            obj.stock = (obj.stock or 0) - cant
            obj.costo_promedio = previo
        elif tipo == "receta":
            obj.producto.receta.remove(obj)
        elif tipo == "tarifa":
            obj.producto.procesos.remove(obj)
        elif tipo == "gasto":
            db.session.delete(obj)
        elif tipo == "activo":
            obj.eliminado = True
        elif tipo == "producto_nuevo":
            usado = (VentaItem.query.join(Venta).filter(VentaItem.producto_id == oid,
                                                        Venta.id.notin_(nuevos["venta"] or [0])).count()
                     or Produccion.query.filter(Produccion.producto_id == oid, Produccion.anulada.is_(False)).count())
            if usado or obj.stock:
                obj.activo = False
                desactivados += 1
            else:
                for pr in Produccion.query.filter(Produccion.producto_id == oid).all():
                    db.session.delete(pr)
                db.session.delete(obj)
        elif tipo == "insumo_nuevo":
            usado = (CompraInsumo.query.filter_by(insumo_id=oid).count()
                     or ProduccionConsumo.query.filter_by(insumo_id=oid).count()
                     or RecetaItem.query.filter_by(insumo_id=oid).count()
                     or PrepRecetaItem.query.filter_by(componente_id=oid).count())
            if usado or (obj.stock or 0) > 1e-9:
                obj.activo = False
                desactivados += 1
            else:
                db.session.delete(obj)
        elif tipo == "cliente_nuevo":
            if Venta.query.filter(Venta.cliente_id == oid, Venta.id.notin_(nuevos["venta"] or [0])).count():
                obj.activo = False
                desactivados += 1
            else:
                db.session.delete(obj)
        elif tipo == "empleado_nuevo":
            if ProduccionMO.query.filter_by(empleado_id=oid).count() or PagoNomina.query.filter_by(empleado_id=oid).count():
                obj.activo = False
                desactivados += 1
            else:
                db.session.delete(obj)
        db.session.flush()
    if errores:
        raise ErrorNegocio(errores)
    return desactivados


def resumen_txt(res):
    """Líneas legibles del resumen de una importación."""
    def m(x):
        return "$ " + f"{x:,.0f}".replace(",", ".")

    def pl(n, uno, varios):
        return f"{n} {uno if n == 1 else varios}"
    out = []
    for k, titulo, *_ in HOJAS:
        d = res.get(k)
        if not d:
            continue
        partes = []
        if d.get("nuevos"):
            partes.append(pl(d["nuevos"], "nuevo", "nuevos"))
        if d.get("actualizados"):
            partes.append(pl(d["actualizados"], "actualizado", "actualizados"))
        if d.get("valor_bodega"):
            partes.append(f"bodega por {m(d['valor_bodega'])}")
        if d.get("referencias"):
            partes.append(f"ficha de {pl(d['referencias'], 'referencia', 'referencias')} ({pl(d.get('lineas', 0), 'material', 'materiales')})")
        if d.get("tarifas"):
            partes.append(pl(d["tarifas"], "tarifa", "tarifas"))
        if d.get("pares"):
            partes.append(f"{pl(d['pares'], 'par', 'pares')} por {m(d.get('valor', 0))}")
        if d.get("facturas"):
            partes.append(f"{pl(d['facturas'], 'factura', 'facturas')} por {m(d.get('valor', 0))}")
        if d.get("clientes_creados"):
            partes.append(pl(d["clientes_creados"], "cliente creado", "clientes creados"))
        if d.get("ventas"):
            partes.append(f"{pl(d['ventas'], 'venta', 'ventas')} por {m(d.get('valor', 0))}")
        if d.get("gastos"):
            partes.append(f"{pl(d['gastos'], 'gasto', 'gastos')} por {m(d.get('valor', 0))}")
        if d.get("maquinas"):
            partes.append(pl(d["maquinas"], "máquina", "máquinas"))
        if d.get("repetidos"):
            partes.append(f"{pl(d['repetidos'], 'ya estaba', 'ya estaban')} (se omiten)")
        if partes:
            out.append((titulo, ", ".join(partes)))
    return out


def a_json(x):
    return json.dumps(x, default=str)
